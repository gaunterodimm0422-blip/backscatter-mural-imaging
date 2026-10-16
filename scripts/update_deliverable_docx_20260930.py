"""Update the two retained explanations from the frozen D2-15% results.

Never modifies the two source DOCX files or any simulation/result file.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "m3/outputs/d2f15"
OUT = ROOT / "deliverables"


def read(name):
    return json.loads((DATA / name).read_text(encoding="utf-8"))


def fit(name):
    j = read(f"{name}/m2_edge_result.json")
    assert all("response_d2f15" in p for p in j["inputs"].values())
    s = j["fits"]["signal"]
    h = j["head_x_mm"]
    return {**s, "world": s["xe"] + h,
            "lo": s["xe_lo"] + h, "hi": s["xe_hi"] + h,
            "bias": s["xe"] + h - j["truth_edge_world_mm"],
            "truth": j["truth_edge_world_mm"], "head": h,
            "null": j["fits"]["null"]["T"],
            "split": j["fits"]["null_split"]["T"],
            "cross": j["fits"]["consistency"]["xe"] - s["xe"]}


M3 = [fit(n) for n in ("m3_edge_result", "m3_offset", "m3_offset2")]
M2 = [fit(n) for n in ("m2_mural_edge", "m2_mural_edge_offset",
                      "m2_mural_edge_offset2", "m2_mural_edge_offset3")]
F3 = read("f3_thickness/f3_summary.json")
C3 = read("m3_d2f15_summary.json")["calibration"]
C2 = read("m2_d2f15_summary.json")["calibration"]


def text(p, value):
    """Replace text without discarding paragraph settings or adjacent figures."""
    assert not p._p.xpath(".//w:drawing"), "Cannot replace an image paragraph"
    p.clear()
    p.add_run(value)


def before(anchor, value, style="Normal"):
    p = anchor.insert_paragraph_before(value, style=style)
    return p


def table_style(t, widths=None, font=10):
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths is None:
        widths = [6.4 / len(t.columns)] * len(t.columns)
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    pr = t._tbl.tblPr
    borders = pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "D9D9D9")
        borders.append(e)
    margins = pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        pr.append(margins)
    for edge in ("top", "bottom", "left", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), "90")
        e.set(qn("w:type"), "dxa")
        margins.append(e)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        for e in list(trpr):
            if e.tag == qn("w:trHeight"):
                trpr.remove(e)
        no_split = OxmlElement("w:cantSplit")
        trpr.append(no_split)
        if ri == 0:
            repeat = OxmlElement("w:tblHeader")
            trpr.append(repeat)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            old = tcpr.find(qn("w:shd"))
            if old is not None:
                tcpr.remove(old)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "E9EEF4" if ri == 0 else "FFFFFF")
            tcpr.append(shd)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                p.paragraph_format.keep_with_next = False
                for r in p.runs:
                    r.font.size = Pt(font)
                    r.font.color.rgb = RGBColor(0, 0, 0)
                    r.bold = ri == 0


def set_table(t, rows, widths=None, font=10):
    assert len(rows[0]) == len(t.columns)
    while len(t.rows) < len(rows):
        t.add_row()
    while len(t.rows) > len(rows):
        tr = t.rows[-1]._tr
        tr.getparent().remove(tr)
    for row, values in zip(t.rows, rows):
        assert len(values) == len(row.cells)
        for cell, value in zip(row.cells, values):
            cell.text = str(value)
    table_style(t, widths, font)


def new_table(doc, anchor, rows, widths=None, font=10):
    t = doc.add_table(rows=1, cols=len(rows[0]))
    anchor._p.addprevious(t._tbl)
    set_table(t, rows, widths, font)
    before(anchor, "")
    return t


def clean_styles(doc):
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        if name not in doc.styles:
            continue
        st = doc.styles[name]
        st.font.color.rgb = RGBColor(0, 0, 0)
        pr = st.element.find(qn("w:pPr"))
        if pr is not None:
            for el in list(pr):
                if el.tag == qn("w:pBdr"):
                    pr.remove(el)
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading") or p.style.name == "Title":
            for r in p.runs:
                r.font.color.rgb = RGBColor(0, 0, 0)
        if p.style.name == "Title":
            pr = p._p.pPr
            if pr is not None:
                for el in list(pr):
                    if el.tag == qn("w:pBdr"):
                        pr.remove(el)
    for t in doc.tables:
        table_style(t)
    doc.core_properties.modified = __import__("datetime").datetime(2026, 9, 30, 8, 40)


def f3_block(doc, anchor, detailed):
    before(anchor, "8.5 界面空鼓厚度扫描" if detailed else "第四，界面空鼓厚度扫描已经完成。",
           "Heading 3" if detailed else "Normal")
    before(anchor, "在 M3 构型下扫描 20、10、5、3、2、1 mm 厚的空气层。20 mm 沿用主档，其他五档各运行400万初级事例；全部按 D1 8%、D2 15% 的响应分析。空气层下表面固定在深40 mm的地仗与支撑体界面，因此越薄的空气层上表面越深，厚度与埋深没有独立分离。")
    before(anchor, "检出要求 T≥30 且至少一个深带的拟合对比度为正。定位还要求边界偏差绝对值≤2 mm、68%区间半宽≤2 mm；换参照场后的边界差须≤5 mm。这些条件在本轮数据产生前冻结。")
    if detailed:
        rows = [["厚度/mm", "T", "偏差/mm", "68%半宽/mm", "换参照场差/mm", "检出", "定位"]]
        for r in F3["rows"]:
            rows.append([r["t_mm"], f'{r["T"]:.1f}', f'{r["bias"]:+.2f}',
                         f'{r["half_width"]:.2f}', f'{r["xe_cross_diff"]:+.2f}',
                         "通过" if r["F3_1"] else "未通过", "通过" if r["located"] else "未通过"])
        new_table(doc, anchor, rows, [0.65, 0.65, 0.85, 1.05, 1.35, 0.85, 0.85], 9.5)
    else:
        rows = [["厚度/mm", "空气层深度/mm", "T", "检出", "定位"]]
        for r in F3["rows"]:
            rows.append([r["t_mm"], f'{40-r["t_mm"]}–40', f'{r["T"]:.1f}',
                         "通过" if r["F3_1"] else "未通过", "通过" if r["located"] else "未通过"])
        new_table(doc, anchor, rows, [0.8, 1.8, 0.9, 1.2, 1.2], 10.5)
    before(anchor, "10 mm档已经能检出（T=87.5），但边界区间半宽4.02 mm，未满足2 mm定位要求；20 mm档偏差−0.07 mm、半宽0.33 mm，定位通过。5 mm及以下四档均未检出，不能把其噪声拟合出来的小区间解释为准确定位。")
    before(anchor, "因此，在本次构型、材料模型、深40 mm界面及400万事例条件下，已扫描档位中最薄的检出通过档为10 mm，最薄的定位通过档为20 mm。每个目标档目前只有一个随机实现，且共享参照场；这不是任意深度下的通用厚度极限，也不是横向空间FWHM，更不是经过重复试验测得的检出概率曲线。")


def algorithm():
    doc = Document(OUT / "成像算法说明_20260912.docx")
    p, t = list(doc.paragraphs), list(doc.tables)
    changes = {
        0: "双511 keV符合标记背散射成像算法说明",
        1: "壁画界面空鼓横向边界定位　更新版　2026-09-30　D1 8%与D2 15% FWHM @511 keV",
        3: "本文说明事件合成、探测器响应、筛选、散射点重建及边界判定。正式结果以M3为主、M2为对照；当前厚度扫描最薄检出通过档为10 mm，最薄定位通过档为20 mm。以下先明确两台探测器的分工。",
        14: "按(runID, eventID)及探测器分组，沉积能量求和，位置按能量加权求质心。当前是模拟事例级合成，不是已经实现的实验脉冲合成。",
        17: "2.1 D2位置读出口径与未实现的改进",
        18: "历史D2位置读出诊断记录了多点相互作用及不同位置口径的差异。下表保留该诊断的数值，但不把它们当作当前M3、D2 15%条件下已经重新验证的空间分辨率。",
        20: "当前流程实际使用能量加权质心。康普顿序列读出的2.53 mm是历史诊断估计，完整读出算法尚未实现；它是候选改进，不是本轮已交付性能。",
        21: "最早时间晶体的1.46 mm依赖理想相互作用顺序，只能作为历史诊断中的理想基准。本项目没有实现足以给晶体内相互作用定序的时间读出，不能将其报为系统分辨率。",
        26: "能量展宽采用高斯模型，FWHM(E)=f×511×√(E/511) keV，σ(E)=FWHM(E)/2.355。D1取f=0.08，D2取f=0.15；D1在255.5 keV处的FWHM约28.9 keV。能量误差经散射角反解放大，是深度误差的重要来源；位置读出、几何条件和多次散射也会影响结果。",
        29: "下列能量、位置、时间差和运动学条件使用施加探测器响应后的读数。但配对仍按模拟事例编号进行，并要求不同初级光子的模拟身份掩码不重叠。这是理想符合条件，不应把整个链条称为完全不使用模拟标签。",
        31: "能窗门（D1）：180–250 keV。按理想单次康普顿关系，下限180 keV对应θ≈147.0°，上限250 keV对应θ≈92.5°，即约92.5°–147.0°。能量展宽会使真实散射角在窗边缘混入或漏出，所以不是对真实角度的硬切。",
        32: "其他门：时间差绝对值≤10 ns；D1有效半径12.5–100 mm；D2有效面半宽25 mm；标记方向极角0°–35°；散射角反解合法。闭式求解还要求θ>α、sinθ非退化，路程在1–280 mm内。当前求解器不要求散射点必须落在真值空气层内。",
        38: "入射光子的估计直线由源位置和D2测量位置确定；它受D2位置读出和理想背靠背假设的影响，并不是无误差的真实轨迹。",
        41: "散射点在估计入射直线上，并满足由D1能量推得的散射角约束。等价地，候选点位于以H为顶点、以−v为轴、半角θ的锥面上。模型假设物体内发生单次康普顿散射，且D1读数可近似代表散射后光子的能量；实际数据中的部分沉积和多次散射会破坏这一假设。",
        49: "v = −(T − S) / |T − S|　（由D2标记反推入射单位方向）",
        53: "在能量和几何合法、θ>α的条件下直接求出物理解，无需迭代。无解或退化事例仍须分类记录。",
        59: "历史C++/ROOT自检中，4000个测试事例有2088个有解，两种解法的最大差值显示为0.000000 mm（输出精度范围内）。这验证两种数学实现的一致性，不代表实际成像误差为零。",
        67: "当前正式Python求解器使用以下数字失败码。旧C++二次方程实现中的额外失败码不与这套编号混用。",
        68: "物体深度分带属于后续分析，不应当用真值缺陷边界裁剪重建点来使成像变好。当前phantom_truth_policy为ignore，不以物体内单次散射真值筛选正式样本；模拟光子身份和事例编号的配对例外已在第四节说明。",
        69: "六、横向误差与深度误差的不同来源",
        70: "横向位置对标记几何较敏感，深度对散射角与能量读数较敏感；两者并非完全独立。当前研究以空鼓的横向边界定位为目标，不声称已经准确恢复空气层的厚度或深度。",
        71: "对cosθ=2−511/E′线性求导，得到dθ/dE′=−511/(E′²sinθ)。下表按D1 8% @511 keV及√E展宽模型估算局部角度标准差。θ接近180°时导数发散，线性传播不再可靠；不能将539°当作实际可定义的角度误差。",
        73: "历史端到端合成测试记录了横向绝对误差中位1.91 mm、深度绝对误差中位11.38 mm，约相差6倍。这些是历史测试的误差中位数，不是FWHM，也不是当前M3的边界定位区间。",
        74: "当前15%响应诊断中，M2与M3的深度残差中央68%半宽分别为15.09和10.93 mm，即(p84−p16)/2；深度绝对误差中位分别为8.81和6.30 mm。改善与几何杠杆臂缩短相符，但本轮同时改变多项几何，不能把全部改善归因于单一参数。这也不能解释为深度完全不可测。",
        78: "对按重建深度分带的横向边界进行离轴标定，边界估计存在向视场中心收缩的现象：M2的k约0.795，M3约0.920。它描述边界估计刻度，不是位置FWHM。",
        79: "7.2 条件选择与深度泄漏",
        80: "条件选择是主要解释：按重建深度分带，真实深度分布受到深度误差及离轴接受几何影响。M3参照场诊断的深度残差中位为−0.15 mm；整体中位接近零，并不排除局部或分带后的偏差。",
        81: "落入某一重建深度带的事例，不一定来自同一真实深度带。离轴后入射与回程的几何接受也会变化，造成边界估计收缩。真值在这里用于机制诊断，不用于正式边界拟合。",
        83: "下表为M2四个边界位置在D2 15%下的结果；M3另有轴心、离轴−8和−16 mm三个点。均按探测头坐标做自由截距、无权线性拟合。各位置共享参照场，拟合参数误差不应解读为独立重复测量验证。",
        85: "M2：k=0.7953±0.0260，b=+0.414±0.255 mm，r²=0.99787；M3：k=0.9197±0.0195，b=+0.017±0.201 mm，r²=0.99955。",
        86: "M2候选模型的描述性比较如下。用拟合系数的误差单位量化距离，不把小样本、共享参照场的拟合当作严格的独立显著性检验。",
        87: "k=0.70来自历史两点外推；正式报告以15%响应下自由截距拟合为准。M3三点标定比M2更接近k=1，但缺少独立验证点，不应由拟合残差直接保证新测量精度。",
        89: "模式(a)将探测头移到边界附近，减少离轴刻度影响；模式(b)固定探测头并用标定系数改正。二者是基于模拟结果提出的作业方案，尚不是现场实验验证结论。",
        90: "下表为M3三个已运行位置、未经刻度改正的正式结果。68%区间是剖面似然给出的边界位置区间，离轴−16 mm档不对称。",
        91: "标定残差尺度s=√(Σ残差²/(n−2))：M2为0.465 mm，M3为0.221 mm。它不是空间FWHM，也不能作为独立新测量的误差上限。",
        92: "模式(b)的新测量误差需同时考虑边界统计不确定性、k与b及其协方差、参考场和模型失配。现有三点M3标定没有形成独立验证后的完整误差预算，因此不沿用历史旧参数预算来作性能保证。",
        94: "8.1 历史MLEM诊断与当前成像图",
        95: "历史合成测试使用过下式的MLEM迭代，用于诊断图像和灵敏度掩模。当前正式M3成像演示使用闭式重建点的二维直方图、空鼓场与参照场的归一化比值及深度带剖面，不是下面这套MLEM输出。",
        98: "历史MLEM测试曾出现低灵敏度边缘放大噪声，使用s_j<0.05×max(s)的掩模后，真值横向线带内图像质量占比为67.5%（均匀基准20%）。这是测试诊断，不是当前M3性能指标。",
        99: "同一历史测试的真值深度带内占比为26.2%（均匀基准25%）。它说明该测试中的深度图像恢复有限，不能据此证明所有几何和所有算法下深度都不可测。",
        100: "当前图像负责展示计数分布，正式结论来自下一节的边界似然拟合。图上的真值轮廓只用于比对，不进入拟合；图上描述性双比值对比度与拟合c_b不是同一个量，正文数字引用后者。",
        110: "条件二项似然可消去每个箱的公共总强度；不同初级事例数用已知曝光比例s归一。它不保证任意材料、几何漂移或偶然符合本底都会相消。",
        113: "无边界时x_e和w不可辨识，因此不直接采用Wilks渐近χ²给p值。阈值沿用见数据前冻结的3000次无边界泊松抽样标定，数值如下；这些抽样次数不是GATE独立重复批次数。",
        114: "阈值取T_thr=30。超过有限次标定的最大值，不能自动解释为已经验证了某个严格的现场误报率。",
        115: "8.3 M2对照组结果",
        116: "下表由m3/outputs/d2f15/m2_mural_edge/m2_edge_result.json的15%响应结果更新。M3主结果见8.4节。",
        117: "M2负对照T=4.20，参照場随机对半T=17.74，均小于30。它们支持流程在已测试无缺陷数据上未触发检出，不能证明分析流程在任何数据上都不会误报。",
        118: "M2换参照场的边界差为+0.43 mm。M3轴心、离轴−8、离轴−16 mm的交叉差分别为−0.29、+0.08、−0.35 mm，满足冻结的5 mm要求。",
        119: "九、M3实际结果与简化预测的失效",
        120: "下表比较M2与M3正式拟合对比度，二者均使用D2 15%。此前的事例级简化蒙卡不是正式性能数据，不能继续把预测值写成新构型的结果。",
        121: "实际有效视场按每箱参照计数≥30确定：M2最深25–35 mm带为±34.5 mm，M3为±31.5 mm；浅带仍为±16.5 mm。照射锥在真深30 mm处半宽33.3 mm是几何覆盖，不等于重建分析中的有效视场。",
        122: "简化预测器在选出率、浅带对比度、视场、四带对比度、泄漏份额和k变化上均有已记录偏差。它只能辅助方案排序。M3的6a（深带对比度）、7（选出率预期范围）、8b（视场扩大预期）未通过，其余相关条件通过；实际视场收窄是工程取舍，不能只说是预测错误而省略。",
        132: "参考场每百万模拟初级选出M2 19759.75、M3 27507.75，比例1.392；乘相应发射锥立体角后，每真实湮没选出率分别为2.894×10⁻³与2.487×10⁻³，比例0.860。相同活度、理想符合条件下，达到相同选出计数约需1.16倍时间；这不是相同定位性能的时间保证，也未包含实际死时间和偶然符合。",
        134: "十一、尚未完成的工作与已完成的厚度扫描",
    }
    for i, s in changes.items():
        text(p[i], s)
    t[0].cell(3, 1).text = "背散射返回光子；单次90°–180°运动学能量170.3–255.5 keV，实际能窗180–250 keV"
    t[0].cell(6, 1).text = "1.55 / 1.55 / 5.0 mm（当前按读出体素解释）"
    t[0].cell(6, 2).text = "1.0 / 1.0 / 2.5 mm（当前按读出体素解释）"
    t[1].cell(1, 2).text = "按(runID,eventID)聚合的沉积能量与位置"
    t[1].cell(5, 2).text = "直方图与比值演示；边界位置x_e及检验量T"
    t[1].cell(5, 3).text = "深度分带、横向分箱、参照计数掩模"
    t[3].cell(0, 1).text = "历史横向诊断值"
    t[3].cell(1, 3).text = "当前默认；4.15 mm非本轮新测值"
    t[3].cell(3, 3).text = "历史估计；完整算法未实现"
    t[4].cell(2, 2).text = "1.55 / 1.55 / 5.0 mm"
    set_table(t[6], [["失败码", "当前Python求解器含义"],
                    ["0", "重建合法"], ["1", "D2与源的方向退化"],
                    ["2", "能量不满足单次康普顿运动学"],
                    ["3", "θ≤α或sinθ退化等几何条件失败"],
                    ["4", "解出的路程不在1–280 mm内"]], [0.8, 5.6])
    t[7].cell(9, 2).text = "发散"
    t[7].cell(9, 3).text = "线性传播失效"
    rows = [["档位", "真边界/mm", "离轴/mm", "重建/mm", "68%区间/mm", "偏差/mm", "T"]]
    for name, f in zip(["轴心", "−8", "−16", "+8"], M2):
        rows.append([name, f'{f["truth"]:.0f}', f'{f["truth"]-f["head"]:.0f}',
                     f'{f["world"]:.2f}', f'[{f["lo"]:.2f}, {f["hi"]:.2f}]',
                     f'{f["bias"]:+.2f}', f'{f["T"]:.1f}'])
    set_table(t[8], rows, [0.6, 0.85, 0.65, 0.8, 1.85, 0.85, 0.7], 9.5)
    set_table(t[9], [["候选模型", "预期", "15%结果的描述性距离"],
                    ["k=1.00", "无收缩", "与M2拟合k相距7.9个拟合标准误"],
                    ["固定平移+2.2 mm", "不随离轴变化", "+8 mm档与该模型相距约5.0个区间半宽"],
                    ["k=0.70", "历史两点外推", "与M2拟合k相距3.7个拟合标准误"],
                    ["k=0.83", "几何预测", "与M2拟合k相距1.3个拟合标准误"]], [1.5, 1.7, 3.2])
    t[10].cell(1, 2).text = "M3轴心偏差−0.07 mm；68%区间半宽0.33 mm"
    t[10].cell(1, 3).text = "拟议精定位模式；尚无现场验证"
    t[10].cell(2, 2).text = "须传播统计与标定不确定性；未独立验证"
    rows = [["离轴/mm", "重建世界x/mm", "偏差/mm", "68%区间/mm", "半宽/mm"]]
    for f in M3:
        rows.append([f'{f["truth"]-f["head"]:.0f}', f'{f["world"]:.2f}',
                     f'{f["bias"]:+.2f}', f'[{f["lo"]:.2f}, {f["hi"]:.2f}]', f'{f["xe_sigma"]:.2f}'])
    # Reuse the five-column original budget table.
    set_table(t[11], rows, [0.75, 1.25, 0.9, 2.35, 1.15])
    f = M2[0]
    set_table(t[13], [["编号", "判据", "D2 15%结果", "结论"],
        ["1", "T≥30", f'T={f["T"]:.1f}', "通过"],
        ["2", "边界偏差绝对值≤10 mm", f'偏差{f["bias"]:+.2f} mm', "通过"],
        ["3", "负对照和对半T<30", f'{f["null"]:.2f} / {f["split"]:.2f}', "通过"],
        ["4", "换参照场边界差≤5 mm", f'{f["cross"]:+.2f} mm', "通过"],
        ["5", "四带拟合c_b全正", "11.99 / 13.85 / 31.91 / 26.51 %", "通过"]], [0.5, 2.0, 3.1, 0.8])
    set_table(t[14], [["重建深度带/mm", "M2拟合c_b", "M3拟合c_b", "变化/百分点"]] +
              [[band, f'{100*a:.2f}%', f'{100*b:.2f}%', f'{100*(b-a):+.2f}']
               for band, a, b in zip(["0–8", "8–15", "15–25", "25–35"], M2[0]["c"], M3[0]["c"])])
    t[15].cell(14, 1).text = "180–250 keV"
    t[15].cell(14, 2).text = "180–250 keV"
    t[15].cell(6, 3).text = "M3冻结几何"
    t[15].cell(7, 3).text = "M3冻结几何"
    t[16].cell(5, 1).text = "已完成6档；最薄检出通过10 mm，定位通过20 mm"
    t[16].cell(5, 2).text = "限于当前界面深度、材料及400万事例条件"
    t[16].cell(5, 3).text = "不自动追加新扫描"
    t[16].cell(6, 1).text = "当前质心读出；序列算法完整实现缺"
    t[16].cell(6, 2).text = "历史2.53 mm估计不是已验证新性能"
    t[16].cell(6, 3).text = "候选改进，需另行设计和验证"
    before(p[119], "8.4 M3主结果与负对照", "Heading 3")
    before(p[119], "M3轴心20 mm空鼓：T=630.1，边界x=−15.07 mm，偏差−0.07 mm，68%区间[−15.40,−14.74] mm，半宽0.33 mm；四带拟合c_b为2.45%、12.45%、31.53%、36.25%。这里的0.33 mm是一个边界参数的统计区间半宽，不是系统空间FWHM。过渡模型参数w=1.56 mm也不能不经转换直接当作FWHM。")
    before(p[119], "M3独立参照场负对照T=25.40、参照场对半T=8.35，均未超过30。但25.40高于本轮2000个纯泊松零样本的最大值23.44（99%分位20.05）。运行宏及主要参数核查无差异，负对照拟合的四带对比度正负混杂；原因仍不能仅凭这些检查证明为纯统计涨落。结论保留为负对照通过、但偏尾。")
    f3_block(doc, p[119], True)
    before(p[24], "实现边界：事件编号与不同初级光子身份用于理想配对，尚无1 ns时间响应和跨事例偶然符合模型；位置参数按体素中心量化，不是已标定的位置高斯FWHM。物体相互作用真值仅用于诊断，正式phantom_truth_policy=ignore。")
    before(p[124], "晶体质量由12.37降至2.05 kg，减少约83.4%，约为原来的1/6，而不是降低一个数量级。")
    doc.add_heading("十二、数据来源与适用范围", level=2)
    doc.add_paragraph("M3与M2正式数字取自m3/outputs/d2f15/下的m2_edge_result.json、m3_d2f15_summary.json和m2_d2f15_summary.json；F3取自f3_thickness/f3_summary.json。数字汇总为numbers_d2f15.json。D2响应随机种子24080601，能量和位置响应在分析时施加，未重跑GATE输运。")
    doc.add_paragraph("体模细泥层5 mm、粗泥层35 mm、支撑体100 mm，密度1.70和2.20 g/cm³为仿真取值而非实测。源按理想背靠背511 keV光子模拟，有限活性区、封装与完整核素衰变未建模；M3配置记录GATE 9.4.2/Geant4 11.4.0，旧数据记录9.4.1/11.3.0，跨构型绝对量比较需说明版本差异。本文尚无实验数据或现场剂量与采集时间验证。")
    clean_styles(doc)
    dest = OUT / "成像算法说明_20260930.docx"
    doc.save(dest)
    return dest


def reviewer():
    doc = Document(OUT / "仿真论文框架说明_给审阅人_20260914.docx")
    p, t = list(doc.paragraphs), list(doc.tables)
    changes = {
        1: "双511 keV符合标记背散射成像　壁画界面空鼓横向边界定位　2026-09-30更新",
        3: "本项目研究能否仅从壁画正面，利用双511 keV符合标记及背散射测量判断界面脱空，并估计其横向边界。本文按论文框架说明物理原理、仿真处理、统计判定和现阶段结果，便于直接向导师汇报。",
        4: "当前主构型为M3、M2为对照，正式响应参数为D1 8%和D2 15% FWHM @511 keV。已有界面厚度扫描：400万事例条件下，已扫描档位最薄检出通过为10 mm，最薄定位通过为20 mm；实验验证尚未完成。",
        10: "第三段引出符合标记方法，并与既有背散射定位研究比较。已有资料提供CSIRO、LANL等工作的文献线索，但尚不能把历史检索写成全球唯一实验、长期无人研究或专利有效状态的确定结论。本文贡献应落在壁画体模、构型比较、边界判定及厚度扫描的具体证据上，正式投稿前再补齐引用。",
        14: "本方案利用正负电子湮没的一对近似背靠背511 keV光子。当前GATE采用理想backtoback源，不是完整钠-22衰变及封装输运模型。",
        15: "D2是源另一侧的像素化GAGG标记探测器，D1是源前方、贴近壁画的环形GAGG背散射探测器。M3中源z=17.5 mm，D1占z=2.5–12.5 mm，D2占z=42.5–62.5 mm；源与D1孔轴对齐，但源位于D1上方，不在环体厚度内。",
        16: "D2命中位置与已知源位置给出标记光子的估计方向，反方向就是入射壁画的估计直线。实际方向有位置读出误差，所以“已知方向”不是完全无误差。",
        17: "标记提供了逐事件的入射方向约束，再结合D1位置和能量求散射点。普通背散射并非原则上不能定位，但没有这一逐事件约束时，通常需要扫描、准直或其他额外约束；不能把现有方法一概写成只能测计数、不能定位。",
        22: "θ是散射转角，E′是D1读数，单位keV。上述反解以物体内单次康普顿散射、D1测得散射后光子能量为前提；能量展宽、部分能量沉积或多次散射会使重建偏离真实作用点。",
        25: "第三步将入射直线与物理解对应的锥面求交。角度和几何合法、路程为正时得到散射点；无解或退化情况要剔除并计数，不是所有事件都必有一个可靠散射点。",
        32: "历史C++/ROOT自检：4000个测试事例中2088个有解，两种算法最大差值显示为0.000000 mm。这验证了数学实现的一致性，而不是说明实际成像误差为零。",
        34: "第一条约束是避免下行照射锥直接打到D1内壁。M3孔内半径的几何下限为(17.5−2.5)×tan35°=10.50 mm，实际取12.5 mm，余量2.00 mm。该检查针对理想照射锥，不等于已经模拟了完整源封装或消除了所有直接本底。",
        35: "第二条约束是照射抽样范围与D2几何接受匹配。源到D2中面的距离35 mm、半宽25 mm，对应内接圆边缘半角35.5°，接近35°发射锥。35°在当前程序中是锥内抽样设置，不能直接当作真实各向同性源的硬件准直效率。",
        40: "事件合成把同一模拟事例在每台探测器中的沉积求和；响应步骤加入能量高斯展宽及位置体素量化；再做读数筛选与闭式重建。D1读出体素为1.55/1.55/5 mm、D2为1/1/2.5 mm；这些体素尺寸不等同于已标定的位置FWHM。",
        41: "先加响应再筛选，避免用理想能量高估能窗内计数。当前仍用事例编号配对，并通过不同初级光子的模拟身份掩码判选，因此属于理想符合分析，不能称为全链条仅用实验可观测量。时间窗10 ns已施加，但1 ns时间分辨和跨事例偶然符合尚未模拟；相互作用点及散射次数真值不进入正式边界拟合。",
        43: "按D2 15%正式结果，每模拟初级的选出率M3/M2=1.392；折算到每真实湮没为0.860。相同活度、理想符合条件下，得到相同选出计数约需1.16倍时间。这不是达到相同定位精度必需的时间，也不是已测现场采集时间。",
        47: "横向位置受标记几何影响较大，深度受能量反解角度的影响较大；两者均会受几何和多次散射影响，不是完全独立的两个测量通道。",
        49: "历史端到端合成测试中，横向绝对误差中位1.91 mm、深度绝对误差中位11.38 mm，约相差6倍。这些是历史测试的误差中位数，不是FWHM，也不是当前M3边界位置区间。",
        50: "因此现阶段把主目标放在横向边界定位，不把重建深度图解释为已经准确测出了空气层厚度。但仍利用深度分带帮助区分缺陷信号，并非声称深度完全没有信息。",
        51: "当前深度残差中央68%半宽(p84−p16)/2由M2的15.09 mm降到M3的10.93 mm，约缩小27.6%。历史诊断的源到D1命中点距离中位由50.72变39.53 mm，与缩短杠杆臂的机制一致；但M3同时改变多项几何，不能将全部改善单独归给D1贴近表面。此半宽也不是深度FWHM。",
        54: "参照场用于比较同几何、同材料模型下空鼓测量的计数变化。共同的强度因素经过曝光归一后可抵消，几何、材料或背景不匹配则未必相消。在当前条件下空鼓侧出现计数亏损，但不能推广成所有能窗和所有深度的空鼓必然只使计数减少。",
        55: "按重建深度分为0–8、8–15、15–25、25–35 mm四带，各带台阶高度及归一参数可不同，共享边界位置和过渡宽度。模型利用同一物理边界的信息，提高拟合约束；不能无条件保证比任何单带分析都准确。横向分箱3 mm，参照计数不足30的箱不参与相应带的拟合。",
        58: "无边界时边界位置与过渡宽度不可辨识，不能直接用Wilks渐近卡方分布换算p值。阈值30高于有限次零样本的最大值，但这本身不是严格现场误报率保证。",
        59: "M3负对照T=25.40、参照场随机对半T=8.35，均低于30。但25.40高于本轮2000个纯泊松零样本的最大值23.44，属于偏尾结果。宏和主要设置核对无差异，拟合对比度正负混杂，仍不能证明其唯一原因就是纯统计涨落。论文需同时报告通过和偏尾，不再新增第三份参照模拟。",
        60: "阈值与判据在见到相应数据前冻结，15%重分析没有事后调整门槛。这里报告的是仿真判定结果，置信区间的重复覆盖率、现场误报率和实验稳健性还没有全面验证。",
        63: "横向刻度系数描述离轴边界的估计收缩，不是空间分辨率。15%下M2四点k=0.7953±0.0260，M3三点k=0.9197±0.0195；M3未改正的离轴−8与−16 mm偏差分别+0.84和+1.21 mm，M2对应+2.17和+3.79 mm。标定点共享参照场，尚缺独立验证点。",
        64: "浅层拟合对比度由M2的11.99%变为M3的2.45%。空鼓位于20–40 mm，真正浅层单次散射的直接信号应较弱。真值拆分诊断中，M3浅带单次散射的描述性双比值对比度为−0.48±2.00%，与零相容；M2为8.55±2.49%。不能说精确等于零。该结果支持深度泄漏减少的解释，但不是独立的空间分辨率标定。",
        65: "最深25–35 mm带的正式拟合c_b由26.51%增至36.25%。真值拆分对比度、成像图描述性对比度与正式拟合c_b的定义不同，不能互相代换。真值仅用于诊断，正式拟合使用未按散射次数真值筛选的样本。",
        68: "第一，准确恢复空气层深度和厚度尚未实现。历史MLEM合成诊断中，横向真值线带内图像质量占67.5%（均匀基准20%），深度带内26.2%（均匀基准25%），只说明该历史测试的恢复局限。当前M3展示的是闭式点云的二维直方图、空鼓场/参照场比值和分带剖面，不是MLEM结果；正式判定来自共享边界拟合。",
        69: "第二，M3最深带的多次散射份额约48.6%。对这些事件强行套单次散射反解，所得点一般不等于真实作用点。目前没有经过验证的可观测筛选将它们完全排除；本次按散射次数拆分用了真值，属于机制诊断，不能转写成已经可用于实验的筛选条件。",
        70: "第三，简化预测器有六类已记录偏差，只能辅助排序。M3深带对比度条件6a、选出率预期区间7、深带视场扩大条件8b未通过，必须如实保留。尤其实际最深带有效窗M2为±34.5 mm、M3为±31.5 mm，收窄是实际工程取舍，不是只需修改预测文字就可消除的现象。",
        74: "第一段总结仿真证据：M3轴心20 mm空鼓T=630.1、边界偏差−0.07 mm、68%区间半宽0.33 mm；已测离轴−16 mm偏差+1.21 mm。负对照未超过门槛，但存在偏尾。0.33 mm是边界参数的统计不确定性，不代表系统能分开相距0.33 mm的两个目标。",
        75: "第二段总结构型取舍：深度残差中央68%半宽缩小27.6%，横向刻度k由0.795变0.920；D1晶体质量12.37降至2.05 kg，减少83.4%，约为原来的1/6，而不是减少一个数量级。折算每真实湮没的选出率降低约14%，深带有效视场也收窄；不能只报告改善。",
        76: "第三段总结已测能力与下一步：当前界面扫描最薄检出通过10 mm、定位通过20 mm，5 mm及以下未检出。后续由目标空鼓厚度、深度及测量时间要求决定是否补做研究；实验验证、时间响应和偶然符合、真实材料及剂量评估仍缺，不能提前作为已完成成果。",
        81: "正式性能数字来自以下15%分析文件。历史数学和合成自检用于实现核对；其数值不冒充当前M3硬件性能。",
        82: "已有宏、配置、分析包的校验和与随机种子记录见项目操作手册。密度1.70与2.20 g/cm³为仿真取值，源为理想背靠背511 keV；有限活性区、封装、完整核素衰变及实验数据尚未加入。软件配置记录M3为GATE 9.4.2/Geant4 11.4.0，旧数据为9.4.1/11.3.0；比较绝对量时应注明。",
    }
    for i, s in changes.items():
        text(p[i], s)
    t[0].cell(4, 2).text = "线性传播失效，不报539°"
    set_table(t[1], [["指标", "M2对照 D2 15%", "M3主构型 D2 15%"],
        ["检验量T", "325.2", "630.1"],
        ["轴心边界偏差/mm", "−0.13", "−0.07"],
        ["68%区间半宽/mm", "0.51", "0.33"],
        ["深度残差中央68%半宽/mm", "15.09", "10.93"],
        ["横向刻度系数k", "0.795±0.026（4点）", "0.920±0.020（3点）"],
        ["D1晶体质量/kg", "12.37", "2.05"],
        ["每真实湮没选出率", "2.894×10⁻³", "2.487×10⁻³（比值0.860）"]], [2.5, 1.95, 1.95])
    t[2].cell(19, 1).text = "历史测试中的迭代图像算法；当前正式演示不是MLEM"
    for row in t[2].rows:
        if row.cells[0].text == "检验量 T":
            row.cells[1].text = "有边界模型相对无边界模型的似然改善，须结合冻结阈值判读"
        if row.cells[0].text == "杠杆臂":
            row.cells[1].text = "源到D1命中点的距离，影响能量误差到位置误差的传播"
    t[2].add_row().cells[0].text = "边界区间半宽与空间FWHM"
    t[2].rows[-1].cells[1].text = "前者是拟合一个边界参数的不确定性；后者是空间响应分布宽度，两者不能混称"
    set_table(t[3], [["内容", "来源或状态"],
        ["两种闭式实现的一致性", "历史C++/ROOT自检，4000测试事例；非当前性能测量"],
        ["能量到角度的局部误差估计", "D1 8%及√E展宽模型；180°线性传播失效"],
        ["M2正式结果与四点标定", "m3/outputs/d2f15/m2_mural_edge及三个offset目录；m2_d2f15_summary.json"],
        ["M3正式结果与三点标定", "m3/outputs/d2f15/m3_edge_result、m3_offset、m3_offset2；m3_d2f15_summary.json"],
        ["深度误差诊断", "m3/outputs/d2f15/diag_m2_ref、diag_m3_ref"],
        ["单次与多次散射拆解", "m3/outputs/d2f15/diag_ms_m2、diag_ms_m3；真值仅供诊断"],
        ["界面厚度扫描F3", "m3/outputs/d2f15/f3_thickness/f3_summary.json；六个厚度档"],
        ["成像演示与边界图", "outputs/m3_edge_ratio与imaging_demo_fig1、fig2，2026-09-28更新；图上对比度不代替拟合c_b"],
        ["全部正式数字汇总", "m3/outputs/d2f15/numbers_d2f15.json"],
        ["源及响应、材料参数", "configs/m3_edge_main_d2f15_20260923.json；密度非实测，理想源模型"],
        ["文献、实验及现场条件", "文献线索待正式核对；实验数据、偶然符合及剂量验证仍缺"]], [2.0, 4.4])
    f3_block(doc, p[71], False)
    before(p[45], "当前能窗：D1 180–250 keV、D2 400–600 keV；D1名义理想角度范围约92.5°–147.0°。D2从原口径更新为15%时，只重做分析响应，没有重跑GATE理想能量输运。")
    clean_styles(doc)
    dest = OUT / "仿真论文框架说明_给审阅人_20260930.docx"
    doc.save(dest)
    return dest


def main():
    sources = [OUT / "成像算法说明_20260912.docx", OUT / "仿真论文框架说明_给审阅人_20260914.docx"]
    hashes = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in sources}
    files = [algorithm(), reviewer()]
    for p, h in hashes.items():
        assert hashlib.sha256(p.read_bytes()).hexdigest() == h, "Original changed"
    for file in files:
        d = Document(file)
        body = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
        assert "630.1" in body and "87.5" in body and "400万" in body
        assert "180–250" in body and "15%" in body
        for stale in ("326.7", "632.6", "0.926", "175–260", "538.90", "空鼓厚度下限 | 未扫描"):
            assert stale not in body, (file.name, stale)
        print(file.name, file.stat().st_size, "bytes", hashlib.sha256(file.read_bytes()).hexdigest())
    print("Two originals preserved. No simulation, result or response file changed.")


if __name__ == "__main__":
    main()
